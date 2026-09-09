#!/usr/bin/env python3
"""Prepare BUA residential-environment observations from a retained raw release."""
import argparse
import csv
import math
from collections import defaultdict
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook

from csv_io import write_csv
from release_manifest import load as load_manifest
from residential_environment import BOUNDARY_RULE, METHOD_VERSION, percentile_for, score_for, weighted_cutpoints
from prepare_local_transport import read_population


RAW_FILES = {
    'iod2025-underlying-indicators-v2.xlsx', 'wimd2025-physical-environment.csv',
    'wimd2025-housing.csv', 'defra-no2-2024.csv', 'defra-pm25-2024.csv',
    'defra-pm10-2024.csv', 'ons-public-green-space-corrected.xlsx',
    'oa21-population-weighted-centroids.csv', 'oa21-lsoa21-msoa21-lookup.csv',
    'oa11-oa21-change-lookup.csv', 'oa11-lsoa11-msoa11-lookup.csv',
    'census2011-ks101ew-oa.zip',
}
PILLARS = ('air', 'quiet', 'green', 'housing_environment')
OUTPUT_COLUMNS = (
    'location_id', 'air_burden', 'no2_ug_m3', 'pm25_ug_m3', 'pm10_ug_m3',
    'noise_exposed_pct', 'green_within_300m_pct', 'green_area_within_1000m_m2',
    'epc_sap_mean', 'air_percentile', 'quiet_percentile', 'green_percentile',
    'housing_environment_percentile', 'environment_index_0_100', 'national_percentile',
    'score', 'air_period', 'quiet_period', 'green_period', 'housing_environment_period',
    'air_evidence_id', 'quiet_evidence_id', 'green_evidence_id', 'housing_environment_evidence_id',
    'geography_code', 'geography_vintage', 'air_population_covered', 'quiet_population_covered',
    'green_population_covered', 'housing_environment_population_covered', 'population_expected',
    'method_version', 'confidence', 'reason',
)
AUDIT_COLUMNS = ('bua_code', 'bua_name',) + OUTPUT_COLUMNS[1:] + ('exclusion_reason',)


def read_csv(path):
    with Path(path).open(newline='', encoding='utf-8-sig') as source:
        return list(csv.DictReader(source))


def keyed(rows, key, value):
    result = {}
    for row in rows:
        if row[key] in result:
            raise ValueError(f'duplicate {key}: {row[key]}')
        result[row[key]] = value(row)
    return result


def read_defra(path):
    with Path(path).open(newline='', encoding='utf-8-sig') as source:
        reader = csv.reader(source)
        metadata = [next(reader) for _ in range(5)]
        header = next(reader)
        value_index = 3
        if (metadata[1][0] != '2024' or metadata[2][0] != 'annual mean' or
                not metadata[3][0].startswith('ug m-3')):
            raise ValueError(f'{path}: unexpected period, metric or unit')
        result = {}
        for row in reader:
            if len(row) < 4 or row[value_index] == 'MISSING':
                continue
            result[(int(row[1]), int(row[2]))] = float(row[value_index])
    return result


def nearest_grid_value(grid, x, y):
    centre = (math.floor(x / 1000) * 1000 + 500, math.floor(y / 1000) * 1000 + 500)
    if centre in grid:
        return grid[centre], False
    for radius in range(1, 6):
        candidates = [(grid[(centre[0] + dx * 1000, centre[1] + dy * 1000)], dx * dx + dy * dy)
                      for dx in range(-radius, radius + 1) for dy in range(-radius, radius + 1)
                      if (centre[0] + dx * 1000, centre[1] + dy * 1000) in grid]
        if candidates:
            return min(candidates, key=lambda item: item[1])[0], True
    raise ValueError(f'No Defra grid value near {x},{y}')


def read_english_indicators(path):
    ws = load_workbook(path, read_only=True, data_only=True)['IoD25 Living Env Domain']
    rows = ws.iter_rows(values_only=True)
    header = list(next(rows)); positions = {name: header.index(name) for name in header}
    result = {}
    for row in rows:
        code = row[positions['LSOA code (2021)']]
        deprivation = row[positions['Housing energy performance deprivation Score']]
        noise = row[positions['Noise pollution']]
        if code and deprivation is not None and noise is not None:
            result[code] = {'epc': 100 - float(deprivation), 'noise': 100 * float(noise)}
    return result


def read_welsh_indicators(physical_path, housing_path):
    physical = defaultdict(dict)
    for row in read_csv(physical_path):
        code, indicator, raw = row['Area code_reference'], row['Indicator_reference'], row['Data values']
        if code.startswith('W01') and raw and raw != '[x]':
            physical[code][indicator] = float(raw)
    housing = {}
    for row in read_csv(housing_path):
        code, indicator, raw = row['Area code_reference'], row['Indicator_reference'], row['Data values']
        if code.startswith('W01') and indicator == 'epc' and raw and raw != '[x]':
            housing[code] = float(raw)
    result = {}
    for code, values in physical.items():
        if 'nopu' in values and code in housing:
            result[code] = {'noise': values['nopu'], 'epc': housing[code]}
    return result


def read_green(path):
    ws = load_workbook(path, read_only=True, data_only=True)['LSOA Parks and Playing Fields']
    rows = ws.iter_rows(values_only=True); header = list(next(rows))
    code_i, area_i, built_i, near_i = 8, 15, 16, 17
    result = {}
    for row in rows:
        if row[code_i] and row[area_i] is not None and row[built_i] not in (None, 0) and row[near_i] is not None:
            result[row[code_i]] = {'proximity': 100 * float(row[near_i]) / float(row[built_i]),
                                   'provision': float(row[area_i])}
    return result


def read_oa11_population(path):
    member = 'ks101ew_2011oa/KS101EWDATA.CSV'
    with ZipFile(path) as archive, archive.open(member) as raw:
        import io
        return {row['GeographyCode']: int(row['KS101EW0001']) for row in csv.DictReader(io.TextIOWrapper(raw, encoding='utf-8-sig'))
                if row['GeographyCode'].startswith(('E00', 'W00'))}


def weighted(values, populations):
    pairs = [(values[code], populations[code]) for code in values if code in populations and populations[code] > 0]
    return sum(value * population for value, population in pairs) / sum(population for _, population in pairs) if pairs else None


def prepare(raw_dir, shared_release_dir, inputs_dir):
    if set(load_manifest(raw_dir)) != RAW_FILES:
        raise ValueError('Residential-environment manifest does not match the required release')
    shared = load_manifest(shared_release_dir)
    if set(shared) != {'connectivity_metrics_2025.ods', 'oa21_bua24_best_fit.csv', 'census2021-ts001.zip'}:
        raise ValueError('Shared local-transport release is invalid')
    population = read_population(shared_release_dir / 'census2021-ts001.zip')
    oa_bua_rows = read_csv(shared_release_dir / 'oa21_bua24_best_fit.csv')
    oa_to_bua = {row['OA21CD']: row['BUA24CD'] for row in oa_bua_rows}; bua_names = {row['BUA24CD']: row['BUA24NM'] for row in oa_bua_rows if row['BUA24CD']}
    centroids = keyed(read_csv(raw_dir / 'oa21-population-weighted-centroids.csv'), 'OA21CD', lambda row: (float(row['X']), float(row['Y'])))
    oa_lsoa21 = keyed(read_csv(raw_dir / 'oa21-lsoa21-msoa21-lookup.csv'), 'OA21CD', lambda row: row['LSOA21CD'])
    english = read_english_indicators(raw_dir / 'iod2025-underlying-indicators-v2.xlsx')
    welsh = read_welsh_indicators(raw_dir / 'wimd2025-physical-environment.csv', raw_dir / 'wimd2025-housing.csv')
    lsoa_indicators = dict(english, **welsh)
    green_lsoa = read_green(raw_dir / 'ons-public-green-space-corrected.xlsx')
    oa11_population = read_oa11_population(raw_dir / 'census2011-ks101ew-oa.zip')
    oa11_lsoa = keyed(read_csv(raw_dir / 'oa11-lsoa11-msoa11-lookup.csv'), 'OA11CD', lambda row: row['LSOA11CD'])
    oa21_oa11 = defaultdict(set)
    for row in read_csv(raw_dir / 'oa11-oa21-change-lookup.csv'):
        if row['OA11CD'] and row['OA21CD']:
            oa21_oa11[row['OA21CD']].add(row['OA11CD'])
    grids = {name: read_defra(raw_dir / filename) for name, filename in
             [('no2', 'defra-no2-2024.csv'), ('pm25', 'defra-pm25-2024.csv'), ('pm10', 'defra-pm10-2024.csv')]}

    oa_values, air_fallbacks = {}, 0
    for oa, pop in population.items():
        if oa not in oa_to_bua or not oa_to_bua[oa] or oa not in centroids or oa not in oa_lsoa21:
            continue
        x, y = centroids[oa]; values = {}
        for pollutant, grid in grids.items():
            values[pollutant], fallback = nearest_grid_value(grid, x, y); air_fallbacks += int(fallback)
        values['air_burden'] = (values['no2'] / 10 + values['pm25'] / 5 + values['pm10'] / 15) / 3
        legacy = [(oa11, oa11_population.get(oa11, 0), green_lsoa.get(oa11_lsoa.get(oa11, '')))
                  for oa11 in oa21_oa11.get(oa, ())]
        # Both 2025 indicator releases identify their observations by 2021 LSOA.
        # Use the exact OA21-to-LSOA21 relationship in both countries.  Routing
        # Welsh observations through predecessor LSOA11 codes loses replacement
        # areas such as Coity Higher 1--4 (W01001981--W01001984).
        values.update(lsoa_indicators.get(oa_lsoa21[oa], {}))
        usable = [(weight, item) for _, weight, item in legacy if weight > 0 and item]
        if usable:
            total = sum(weight for weight, _ in usable)
            values['green_proximity'] = sum(weight * item['proximity'] for weight, item in usable) / total
            values['green_provision'] = sum(weight * item['provision'] for weight, item in usable) / total
        oa_values[oa] = values

    bua_oas = defaultdict(list)
    for oa, bua in oa_to_bua.items():
        if bua and oa in population:
            bua_oas[bua].append(oa)
    raw_buas = {}
    pillar_fields = {'air': ('air_burden', 'no2', 'pm25', 'pm10'), 'quiet': ('noise',),
                     'green': ('green_proximity', 'green_provision'), 'housing_environment': ('epc',)}
    for bua, oas in bua_oas.items():
        expected = sum(population[oa] for oa in oas)
        row = {'bua_code': bua, 'bua_name': bua_names[bua], 'population_expected': expected}
        for pillar, fields in pillar_fields.items():
            covered = [oa for oa in oas if oa in oa_values and all(field in oa_values[oa] for field in fields)]
            row[f'{pillar}_population_covered'] = sum(population[oa] for oa in covered)
            for field in fields:
                row[field] = weighted({oa: oa_values[oa][field] for oa in covered}, population)
        raw_buas[bua] = row

    complete = {code: row for code, row in raw_buas.items()
                if all(row[f'{pillar}_population_covered'] == row['population_expected'] > 0 for pillar in PILLARS)}
    if not complete:
        raise ValueError('No complete national BUA observations')
    populations_by_bua = {code: row['population_expected'] for code, row in complete.items()}
    distributions = {
        'air': [(row['air_burden'], populations_by_bua[code]) for code, row in complete.items()],
        'quiet': [(row['noise'], populations_by_bua[code]) for code, row in complete.items()],
        'green_proximity': [(row['green_proximity'], populations_by_bua[code]) for code, row in complete.items()],
        'green_provision': [(row['green_provision'], populations_by_bua[code]) for code, row in complete.items()],
        'housing': [(row['epc'], populations_by_bua[code]) for code, row in complete.items()],
    }
    for code, row in complete.items():
        row['air_percentile'] = percentile_for(row['air_burden'], distributions['air'], True)
        row['quiet_percentile'] = percentile_for(row['noise'], distributions['quiet'], True)
        row['green_percentile'] = (percentile_for(row['green_proximity'], distributions['green_proximity']) +
                                   percentile_for(row['green_provision'], distributions['green_provision'])) / 2
        row['housing_environment_percentile'] = percentile_for(row['epc'], distributions['housing'])
        row['environment_index'] = sum(row[f'{pillar}_percentile'] for pillar in PILLARS) / 4
    index_distribution = [(row['environment_index'], populations_by_bua[code]) for code, row in complete.items()]
    cutpoints = weighted_cutpoints(index_distribution)
    for row in complete.values():
        row['national_percentile'] = percentile_for(row['environment_index'], index_distribution)
        row['score'] = score_for(row['environment_index'], cutpoints)

    def format_row(row, ident='', codes=()):
        reason = ('Population-weighted across the reviewed April 2024 built-up area; '
                  'air uses OA population-weighted centroids, small-area indicators use exact-fit lookups, '
                  'and the structural green-space baseline uses constituent 2011 OA population weights.')
        return {
            'location_id': ident, 'air_burden': f"{row['air_burden']:.4f}", 'no2_ug_m3': f"{row['no2']:.3f}",
            'pm25_ug_m3': f"{row['pm25']:.3f}", 'pm10_ug_m3': f"{row['pm10']:.3f}",
            'noise_exposed_pct': f"{row['noise']:.3f}", 'green_within_300m_pct': f"{row['green_proximity']:.3f}",
            'green_area_within_1000m_m2': f"{row['green_provision']:.2f}", 'epc_sap_mean': f"{row['epc']:.3f}",
            'air_percentile': f"{row['air_percentile']:.3f}", 'quiet_percentile': f"{row['quiet_percentile']:.3f}",
            'green_percentile': f"{row['green_percentile']:.3f}",
            'housing_environment_percentile': f"{row['housing_environment_percentile']:.3f}",
            'environment_index_0_100': f"{row['environment_index']:.3f}",
            'national_percentile': f"{row['national_percentile']:.3f}", 'score': str(row['score']),
            'air_period': '2024 annual mean', 'quiet_period': '2021 strategic noise model',
            'green_period': '2020 structural baseline (corrected 2022)', 'housing_environment_period': 'EPCs 2012-2024',
            'air_evidence_id': 'ENV-AIR-DEFRA-2024', 'quiet_evidence_id': 'ENV-NOISE-IOD-WIMD-2025',
            'green_evidence_id': 'ENV-GREEN-ONS-2020', 'housing_environment_evidence_id': 'ENV-EPC-IOD-WIMD-2025',
            'geography_code': ';'.join(codes), 'geography_vintage': 'April 2024',
            **{f'{pillar}_population_covered': str(row[f'{pillar}_population_covered']) for pillar in PILLARS},
            'population_expected': str(row['population_expected']), 'method_version': METHOD_VERSION,
            'confidence': 'Medium', 'reason': reason,
        }

    locations = {row['id'] for row in read_csv(inputs_dir / 'locations.csv')}
    components = defaultdict(list)
    for row in read_csv(inputs_dir / 'location_geography_components.csv'):
        components[row['location_id']].append(row['geography_code'])
    incomplete_locations = []
    for ident in sorted(locations):
        codes = components.get(ident, [])
        if codes and all(code in complete for code in codes):
            continue
        details = []
        for code in codes:
            row = raw_buas.get(code)
            if not row:
                details.append(f'{code}: absent')
            else:
                missing = [f'{pillar}={row[f"{pillar}_population_covered"]}/{row["population_expected"]}'
                           for pillar in PILLARS
                           if row[f'{pillar}_population_covered'] != row['population_expected']]
                details.append(f'{code}: {", ".join(missing) or "not in complete reference"}')
        incomplete_locations.append(f'{ident} ({"; ".join(details) or "no geography components"})')
    if incomplete_locations:
        raise ValueError('Residential environment is incomplete: ' + '; '.join(incomplete_locations))

    output = []
    for ident in sorted(locations):
        codes = components.get(ident, [])
        expected = sum(complete[code]['population_expected'] for code in codes)
        combined = {'population_expected': expected}
        numeric = ('air_burden', 'no2', 'pm25', 'pm10', 'noise', 'green_proximity', 'green_provision', 'epc',
                   'air_percentile', 'quiet_percentile', 'green_percentile', 'housing_environment_percentile',
                   'environment_index', 'national_percentile')
        for field in numeric:
            combined[field] = sum(complete[code][field] * complete[code]['population_expected'] for code in codes) / expected
        for pillar in PILLARS:
            combined[f'{pillar}_population_covered'] = sum(complete[code][f'{pillar}_population_covered'] for code in codes)
        combined['score'] = score_for(combined['environment_index'], cutpoints)
        output.append(format_row(combined, ident, codes))

    audits = []
    for code, row in sorted(raw_buas.items()):
        if code in complete:
            audit = {'bua_code': code, 'bua_name': row['bua_name'], **format_row(row)}
            audit['exclusion_reason'] = ''
        else:
            audit = {'bua_code': code, 'bua_name': row['bua_name'], 'population_expected': row['population_expected'],
                     'exclusion_reason': '; '.join(p for p in PILLARS if row[f'{p}_population_covered'] != row['population_expected'])}
            for pillar in PILLARS:
                audit[f'{pillar}_population_covered'] = row[f'{pillar}_population_covered']
        audits.append(audit)
    release = {
        'release_id': 'residential-environment-2026-09-09', 'method_version': METHOD_VERSION,
        'reference_bua_count': str(len(complete)), 'reference_population': str(sum(populations_by_bua.values())),
        'q20': f'{cutpoints[0]:.3f}', 'q40': f'{cutpoints[1]:.3f}', 'q60': f'{cutpoints[2]:.3f}',
        'q80': f'{cutpoints[3]:.3f}', 'boundary_rule': BOUNDARY_RULE,
        'pillar_weights': 'air=25;quiet=25;green=25;housing_environment=25',
        'air_grid_nearest_fallback_oa_count': str(air_fallbacks),
    }
    return output, audits, release


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--release-dir', type=Path, default=Path('data/raw/environment/2026-09-09'))
    parser.add_argument('--shared-release-dir', type=Path, default=Path('data/raw/releases/2026-09-09-local-transport'))
    parser.add_argument('--inputs-dir', type=Path, default=Path('data/inputs'))
    parser.add_argument('--canonical-output', type=Path, default=Path('data/inputs/residential_environment.csv'))
    parser.add_argument('--release-output', type=Path, default=Path('data/inputs/residential_environment_release.csv'))
    parser.add_argument('--audit-output', type=Path, default=Path('data/derived/residential_environment_bua_audit.csv'))
    args = parser.parse_args()
    rows, audit, release = prepare(args.release_dir, args.shared_release_dir, args.inputs_dir)
    write_csv(args.canonical_output, rows, OUTPUT_COLUMNS)
    write_csv(args.audit_output, audit, AUDIT_COLUMNS)
    write_csv(args.release_output, [release], release)
    print(f'Prepared {len(rows)} candidate locations from {release["reference_bua_count"]} complete BUAs.')


if __name__ == '__main__':
    main()
