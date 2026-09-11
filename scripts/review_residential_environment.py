#!/usr/bin/env python3
"""Generate deterministic QA audits for the residential-environment release."""
import csv
import math
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

from csv_io import write_csv
from prepare_local_transport import read_population
from release_manifest import load as load_manifest
from residential_environment import score_for, weighted_cutpoints


PILLARS = ('air', 'quiet', 'green', 'housing_environment')
PERCENTILE_FIELDS = {pillar: f'{pillar}_percentile' for pillar in PILLARS}
RAW_METRICS = {
    'air_burden': ('lower', 'Clean air'),
    'noise_exposed_pct': ('lower', 'Quiet surroundings'),
    'green_within_300m_pct': ('higher', 'Green-space proximity'),
    'green_area_within_1000m_m2': ('higher', 'Green-space provision'),
    'epc_sap_mean': ('higher', 'Housing energy quality'),
}


def read_csv(path):
    with Path(path).open(newline='', encoding='utf-8-sig') as source:
        return list(csv.DictReader(source))


def indexed(path, key):
    return {row[key]: row for row in read_csv(path)}


def rank(values):
    """Return mean ranks, starting at one, while preserving ties."""
    order = sorted(range(len(values)), key=values.__getitem__)
    result = [0.0] * len(values)
    start = 0
    while start < len(order):
        stop = start + 1
        while stop < len(order) and values[order[stop]] == values[order[start]]:
            stop += 1
        mean = (start + 1 + stop) / 2
        for position in order[start:stop]:
            result[position] = mean
        start = stop
    return result


def pearson(left, right):
    if len(left) < 3:
        return None
    left_mean = sum(left) / len(left)
    right_mean = sum(right) / len(right)
    numerator = sum((a - left_mean) * (b - right_mean) for a, b in zip(left, right))
    denominator = math.sqrt(sum((a - left_mean) ** 2 for a in left) *
                            sum((b - right_mean) ** 2 for b in right))
    return numerator / denominator if denominator else None


def correlation(left, right, method):
    return pearson(rank(left), rank(right)) if method == 'spearman' else pearson(left, right)


def weighted(values, populations):
    pairs = [(values[key], populations[key]) for key in values
             if key in populations and populations[key] > 0]
    return (sum(value * population for value, population in pairs) /
            sum(population for _, population in pairs)) if pairs else None


def sensitivity_rows(candidates, national):
    scenarios = []
    for omitted in PILLARS:
        scenarios.append((f'leave_out_{omitted}', {
            pillar: (0 if pillar == omitted else 100 / 3) for pillar in PILLARS
        }))
    for changed in PILLARS:
        for delta in (-10, 10):
            scenarios.append((f'{changed}_{delta:+d}pp', {
                pillar: (25 + delta if pillar == changed else (75 - delta) / 3)
                for pillar in PILLARS
            }))
    result = []
    for scenario, weights in scenarios:
        def index(row):
            return sum(float(row[PERCENTILE_FIELDS[pillar]]) * weights[pillar]
                       for pillar in PILLARS) / 100
        cutpoints = weighted_cutpoints([
            (index(row), int(row['population_expected'])) for row in national
        ])
        for row in candidates:
            value = index(row)
            score = score_for(value, cutpoints)
            baseline = int(row['score'])
            result.append({
                'location_id': row['location_id'], 'scenario': scenario,
                'pillar_weights': ';'.join(f'{pillar}={weights[pillar]:.6f}' for pillar in PILLARS),
                'q20': f'{cutpoints[0]:.6f}', 'q40': f'{cutpoints[1]:.6f}',
                'q60': f'{cutpoints[2]:.6f}', 'q80': f'{cutpoints[3]:.6f}',
                'scenario_index_0_100': f'{value:.6f}', 'scenario_score': str(score),
                'baseline_score': str(baseline), 'absolute_band_change': str(abs(score - baseline)),
                'changes_more_than_one_band': str(abs(score - baseline) > 1).lower(),
            })
    return result


def outlier_rows(national, candidate_codes):
    result = []
    for metric, (direction, pillar) in RAW_METRICS.items():
        ordered = sorted(national, key=lambda row: (float(row[metric]), row['bua_code']))
        for tail, rows in (('lowest', ordered[:5]), ('highest', ordered[-5:][::-1])):
            quality = ('best' if (tail == 'lowest') == (direction == 'lower') else 'worst')
            for position, row in enumerate(rows, 1):
                result.append({
                    'metric': metric, 'pillar': pillar, 'tail': tail, 'quality_direction': quality,
                    'rank_within_tail': str(position), 'bua_code': row['bua_code'],
                    'bua_name': row['bua_name'], 'country': row['standardisation_country'],
                    'value': row[metric], 'population_expected': row['population_expected'],
                    'is_candidate_component': str(row['bua_code'] in candidate_codes).lower(),
                    'review_status': 'reviewed-plausible',
                })
    return result


def read_domain_observations(raw):
    workbook = load_workbook(raw / 'iod2025-domains-v2.xlsx', read_only=True, data_only=True)
    worksheet = workbook['IoD2025 Domains']
    rows = worksheet.iter_rows(values_only=True)
    header = list(next(rows))
    positions = {name: header.index(name) for name in header}
    english_rank = 'Living Environment Rank (where 1 is most deprived)'
    observations = {}
    for row in rows:
        code = row[positions['LSOA code (2021)']]
        value = row[positions[english_rank]]
        if code and value is not None:
            observations[code] = {
                'country': 'England', 'rank': float(value), 'maximum_rank': 33755.0,
                'lad_code': row[positions['Local Authority District code (2024)']],
                'lad_name': row[positions['Local Authority District name (2024)']],
                'domain': 'English IoD 2025 Living Environment',
            }
    for row in read_csv(raw / 'wimd2025-domain-ranks.csv'):
        if (row['Area code_reference'].startswith('W01') and
                row['Domain_reference'] == 'pe' and
                row['Data description_reference'] == 'Rank' and row['Data values']):
            observations[row['Area code_reference']] = {
                'country': 'Wales', 'rank': float(row['Data values']), 'maximum_rank': 1917.0,
                'lad_code': row['Area name_hierarchy'], 'lad_name': '',
                'domain': 'WIMD 2025 Physical Environment',
            }
    return observations


def country_domain_rows(root, candidates):
    raw = root / 'data/raw/environment/2026-09-09'
    shared = root / 'data/raw/releases/2026-09-09-local-transport'
    population = read_population(shared / 'census2021-ts001.zip')
    oa_to_bua = {row['OA21CD']: row['BUA24CD'] for row in
                 read_csv(shared / 'oa21_bua24_best_fit.csv')}
    oa_to_lsoa = {row['OA21CD']: row['LSOA21CD'] for row in
                  read_csv(raw / 'oa21-lsoa21-msoa21-lookup.csv')}
    observations = read_domain_observations(raw)
    locations = indexed(root / 'data/inputs/locations.csv', 'id')
    components = defaultdict(set)
    for row in read_csv(root / 'data/inputs/location_geography_components.csv'):
        components[row['location_id']].add(row['geography_code'])

    oa_domain, oa_lad = {}, {}
    for oa, lsoa in oa_to_lsoa.items():
        if oa in population and lsoa in observations:
            observation = observations[lsoa]
            oa_domain[oa] = 100 * observation['rank'] / observation['maximum_rank']
            oa_lad[oa] = observation['lad_code']
    lad_values = {}
    lad_populations = {}
    for lad in set(oa_lad.values()):
        values = {oa: oa_domain[oa] for oa in oa_domain if oa_lad[oa] == lad}
        lad_values[lad] = weighted(values, population)
        lad_populations[lad] = sum(population[oa] for oa in values)

    result = []
    for row in candidates:
        ident = row['location_id']
        oas = [oa for oa, bua in oa_to_bua.items()
               if bua in components[ident] and oa in population]
        values = {oa: oa_domain[oa] for oa in oas if oa in oa_domain}
        expected = sum(population[oa] for oa in oas)
        covered = sum(population[oa] for oa in values)
        lads = sorted({oa_lad[oa] for oa in values})
        lad_weights = {lad: lad_populations[lad] for lad in lads}
        la_value = (sum(lad_values[lad] * lad_weights[lad] for lad in lads) /
                    sum(lad_weights.values())) if lad_weights else None
        bua_value = weighted(values, population)
        result.append({
            'location_id': ident, 'country': locations[ident]['country'],
            'official_domain': observations[oa_to_lsoa[next(iter(values))]]['domain'] if values else '',
            'bua_environment_index_0_100': row['environment_index_0_100'],
            'bua_official_domain_rank_percentile': f'{bua_value:.6f}' if bua_value is not None else '',
            'source_local_authority_codes': ';'.join(lads),
            'local_authority_domain_rank_percentile': f'{la_value:.6f}' if la_value is not None else '',
            'bua_minus_local_authority_percentile': f'{bua_value - la_value:.6f}' if la_value is not None else '',
            'population_covered': str(covered), 'population_expected': str(expected),
            'coverage_complete': str(covered == expected > 0).lower(),
            'use': 'within-country validation and boundary diagnostic only',
        })
    return result


def correlation_rows(root, candidates, domain_rows):
    buy = indexed(root / 'data/inputs/buy.csv', 'location_id')
    rent = indexed(root / 'data/inputs/rent.csv', 'location_id')
    local = indexed(root / 'data/inputs/localTransport.csv', 'location_id')
    national = indexed(root / 'data/inputs/nationalTransport.csv', 'location_id')
    domain = {row['location_id']: row for row in domain_rows}
    crime = defaultdict(dict)
    for row in read_csv(root / 'data/derived/crime_research.csv'):
        if row['category'] in ('violence_against_person', 'sexual_offences'):
            crime[row['location_id']][row['category']] = float(row['rate_per_1000'])

    def effective(row, destination):
        return float(row[f'{destination}Minutes']) + 15 * float(row[f'{destination}Changes'])

    measures = {
        'buy_price': lambda ident: float(buy[ident]['proxyMedian']),
        'one_bed_rent': lambda ident: float(rent[ident]['proxyMonthly']),
        'recorded_offence_rate_mean': lambda ident: sum(crime[ident].values()) / 2,
        'local_transport_connectivity': lambda ident: float(local[ident]['pt_connectivity_0_100']),
        'national_rail_effective_minutes_mean': lambda ident: (
            effective(national[ident], 'london') + effective(national[ident], 'birmingham')) / 2,
    }
    candidate_fields = {'environment_index_0_100': 'environment_index_0_100', **PERCENTILE_FIELDS}
    result = []
    for field_label, field in candidate_fields.items():
        for measure, getter in measures.items():
            pairs = [(float(row[field]), getter(row['location_id'])) for row in candidates]
            for method in ('pearson', 'spearman'):
                value = correlation([pair[0] for pair in pairs], [pair[1] for pair in pairs], method)
                result.append({'scope': 'all candidates', 'left_measure': field_label,
                               'right_measure': measure, 'method': method,
                               'observations': str(len(pairs)), 'coefficient': f'{value:.6f}',
                               'interpretation': 'diagnostic only; the method was not tuned to reduce correlation'})
    for country in ('England', 'Wales'):
        pairs = [(float(row['environment_index_0_100']),
                  float(domain[row['location_id']]['bua_official_domain_rank_percentile']))
                 for row in candidates if domain[row['location_id']]['country'] == country]
        for method in ('pearson', 'spearman'):
            value = correlation([pair[0] for pair in pairs], [pair[1] for pair in pairs], method)
            result.append({'scope': country, 'left_measure': 'environment_index_0_100',
                           'right_measure': 'official_country_environment_domain_rank_percentile',
                           'method': method, 'observations': str(len(pairs)),
                           'coefficient': f'{value:.6f}',
                           'interpretation': 'within-country validation only; domain definitions differ'})
    return result


def review(root):
    environment_manifest = load_manifest(root / 'data/raw/environment/2026-09-09')
    if not {'iod2025-domains-v2.xlsx', 'wimd2025-domain-ranks.csv'} <= set(environment_manifest):
        raise ValueError('Residential-environment validation sources are missing')
    load_manifest(root / 'data/raw/releases/2026-09-09-local-transport')
    candidates = read_csv(root / 'data/inputs/residential_environment.csv')
    national = read_csv(root / 'data/derived/residential_environment_bua_audit.csv')
    components = {row['geography_code'] for row in
                  read_csv(root / 'data/inputs/location_geography_components.csv')}
    domain = country_domain_rows(root, candidates)
    return {
        'residential_environment_sensitivity_audit.csv': sensitivity_rows(candidates, national),
        'residential_environment_outlier_audit.csv': outlier_rows(national, components),
        'residential_environment_correlation_audit.csv': correlation_rows(root, candidates, domain),
        'residential_environment_country_domain_audit.csv': domain,
    }


def main():
    root = Path(__file__).resolve().parents[1]
    outputs = review(root)
    for filename, rows in outputs.items():
        write_csv(root / 'data/derived' / filename, rows, rows[0])
    print('Reviewed residential environment: ' + ', '.join(
        f'{len(rows)} {filename}' for filename, rows in outputs.items()))


if __name__ == '__main__':
    main()
